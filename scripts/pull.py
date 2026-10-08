"""Step 1 of 3: fetch the Google Sheet and write one CSV per tab into data/raw/.

The sheet must be published to the web as a Microsoft Excel (.xlsx) file
(File > Share > Publish to web > Entire document > Microsoft Excel). Paste that
link into config/sheet-url.txt. If the file is empty, this step does nothing and
the CSVs already in data/raw/ are used as they are.
"""
import csv, datetime, io, json, pathlib, sys, urllib.request

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
SCHEMA = json.loads((ROOT / "data" / "schema.json").read_text(encoding="utf-8"))
URL_FILE = ROOT / "config" / "sheet-url.txt"


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def main():
    url = ""
    if URL_FILE.exists():
        lines = [l.strip() for l in URL_FILE.read_text(encoding="utf-8").splitlines()]
        url = next((l for l in lines if l.startswith("http")), "")
    if not url:
        print("No sheet link in config/sheet-url.txt yet - using the CSVs already in data/raw/.")
        return 0

    from openpyxl import load_workbook  # installed by the workflow

    print("Downloading the published sheet ...")
    req = urllib.request.Request(url, headers={"User-Agent": "manitoba-resource-decisions-build"})
    with urllib.request.urlopen(req, timeout=120) as r:
        blob = r.read()
    try:
        wb = load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    except Exception:
        print("ERROR: the link did not return an Excel file. Publish the ENTIRE document as "
              "'Microsoft Excel (.xlsx)' and paste that link into config/sheet-url.txt.")
        return 1

    missing = [t for t in SCHEMA["tabs"] if t not in wb.sheetnames]
    if missing:
        print("ERROR: these tabs are missing from the sheet (names must match exactly): " + ", ".join(missing))
        return 1

    RAW.mkdir(parents=True, exist_ok=True)
    for tab in SCHEMA["tabs"]:
        rows = [[cell_text(v) for v in row] for row in wb[tab].iter_rows(values_only=True)]
        # drop fully empty rows and trailing empty columns
        rows = [r for r in rows if any(c != "" for c in r)]
        if rows:
            width = max(i + 1 for r in rows for i, c in enumerate(r) if c != "")
            rows = [r[:width] + [""] * (width - len(r[:width])) for r in rows]
        with open(RAW / f"{tab}.csv", "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh, lineterminator="\n").writerows(rows)
        print(f"  {tab:14} {max(len(rows) - 1, 0):4} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
